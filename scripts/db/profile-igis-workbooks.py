"""Read all source cells without modifying Excel. Outputs stay in ignored private storage."""
import argparse, collections, datetime, decimal, hashlib, json, os, re, shutil, unicodedata, warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore', message='Workbook contains no default style')
parser = argparse.ArgumentParser()
parser.add_argument('source_dir')
parser.add_argument('--out', default='data/private/igis-master-20261006')
args = parser.parse_args()
out = Path(args.out)
out.mkdir(parents=True, exist_ok=True, mode=0o700)
os.chmod(out, 0o700)
(out/'sources').mkdir(exist_ok=True, mode=0o700)

def clean(v):
    if isinstance(v, str):
        v = unicodedata.normalize('NFC', v).replace('\xa0', ' ').strip()
        return v or None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v

def token(v):
    return json.dumps(v, ensure_ascii=False, sort_keys=True, default=str)

def dump(name, obj):
    p = out/name
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str))
    os.chmod(p, 0o600)

def typ(v):
    if v is None: return 'null'
    if isinstance(v, bool): return 'boolean'
    if isinstance(v, int) or isinstance(v, float) and v.is_integer(): return 'integer'
    if isinstance(v, float): return 'decimal'
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}',str(v)):
        try: datetime.date.fromisoformat(v); return 'date_string'
        except ValueError: return 'invalid_date_string'
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}(\.\d+)?',str(v)): return 'timestamp_string'
    if re.fullmatch(r'[+-]?\d+(\.\d+)?',str(v)): return 'numeric_string'
    return 'string'

books = []
for source in sorted(Path(args.source_dir).glob('*20261006.xlsx')):
    if source.name.startswith('~$'): continue
    name = unicodedata.normalize('NFC',source.name)
    raw = source.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    dest = out/'sources'/name
    dest.write_bytes(raw); os.chmod(dest,0o600)
    wb = openpyxl.load_workbook(dest, data_only=False)
    cached = openpyxl.load_workbook(dest, data_only=True)
    book = dict(file=name,source_path=str(source),sha256=sha,bytes=len(raw),sheets=[])
    for ws in wb:
        header_rows = 2 if name.startswith('펀드 AUM') else 1
        groups = {}
        for merged in ws.merged_cells.ranges:
            if merged.min_row == 1:
                for col in range(merged.min_col, merged.max_col+1): groups[col] = clean(ws.cell(1,merged.min_col).value)
        headers=[]
        for col in range(1,ws.max_column+1):
            group=groups.get(col,clean(ws.cell(1,col).value))
            leaf=clean(ws.cell(header_rows,col).value) or group
            label = f'{group} / {leaf}' if header_rows==2 and group and group!=leaf else leaf
            headers.append(label or f'unnamed_{col}')
        assert len(set(headers))==len(headers),(name,'duplicate headers')
        rows=[]; cells=[]; formulas=[]; errors=[]; blank_rows=[]
        for row in ws.iter_rows(min_row=header_rows+1):
            values=[]
            for c in row:
                if c.data_type=='f':
                    formulas.append(dict(cell=c.coordinate,formula=c.value,cached=cached[ws.title][c.coordinate].value))
                    v=cached[ws.title][c.coordinate].value
                else:v=c.value
                values.append(clean(v))
                if c.value is not None:
                    cells.append(dict(cell=c.coordinate,value=c.value,type=c.data_type,format=c.number_format))
                if c.data_type=='e':errors.append(dict(cell=c.coordinate,value=c.value))
            if all(v is None for v in values): blank_rows.append(row[0].row);continue
            rows.append(dict(excel_row=row[0].row,values=dict(zip(headers,values))))
        fields=[]
        for col,h in enumerate(headers,1):
            vals=[r['values'][h] for r in rows]; non=[v for v in vals if v is not None]
            types=collections.Counter(typ(v) for v in vals); freq=collections.Counter(token(v) for v in non)
            fields.append(dict(column=openpyxl.utils.get_column_letter(col),label=h,nonempty=len(non),blank=len(vals)-len(non),types=dict(types),distinct=len(freq),max_length=max([len(str(v)) for v in non] or [0]),examples=[json.loads(k) for k,_ in freq.most_common(6)],numeric_min=min([v for v in non if isinstance(v,(int,float))] or [None]),numeric_max=max([v for v in non if isinstance(v,(int,float))] or [None]),excel_formats=dict(collections.Counter(c['format'] for c in cells if re.sub(r'\d','',c['cell'])==openpyxl.utils.get_column_letter(col)))))
        duplicates=collections.defaultdict(list)
        for r in rows:duplicates[token(r['values'])].append(r['excel_row'])
        sheet=dict(name=ws.title,state=ws.sheet_state,header_rows=header_rows,range=ws.calculate_dimension(),merged=[str(x) for x in ws.merged_cells.ranges],rows=rows,fields=fields,formulas=formulas,errors=errors,blank_rows=blank_rows,exact_duplicate_rows=[v for v in duplicates.values() if len(v)>1],raw_cells=cells)
        book['sheets'].append(sheet)
    books.append(book)
dump('workbooks.json',books)
summary=[dict(file=b['file'],sha256=b['sha256'],sheets=[{k:v for k,v in s.items() if k not in ('rows','raw_cells')} for s in b['sheets']]) for b in books]
dump('profiles.json',summary)
for b in summary:
    for s in b['sheets']:
        print(json.dumps(dict(file=b['file'],sheet=s['name'],rows=sum(s['fields'][0]['types'].values()),columns=len(s['fields']),empty_columns=[f['label'] for f in s['fields'] if not f['nonempty']],formulas=len(s['formulas']),errors=len(s['errors']),duplicate_groups=len(s['exact_duplicate_rows'])),ensure_ascii=False))
